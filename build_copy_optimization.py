#!/usr/bin/env python3
"""Build the Well and Good Growth site-v2 copy optimization deck as a .docx.

For every discrete chunk of visible copy in site-v2, the deck shows:
  A. Original   the current copy
  B. Optimized  the recommended copy

Method: each page was run through the copy-optimization skill stack in order,
cro (structure and conversion) then angles (emotional framing) then
copywriting-101 (the words, with a strict fact guardrail) then
sentence-conscious-writing (final read-aloud polish). Where a chunk was already
strong, it is listed as reviewed and kept, in the spirit of changing only what
has a real problem.

Keep this script em-dash free. Re-run to regenerate the .docx after edits.
"""
import pathlib

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

ACCENT = RGBColor(0x1A, 0x53, 0x4F)   # deep teal, B label and headings
MUTED = RGBColor(0x8A, 0x8A, 0x8A)    # gray, A label and original text
GRAYNOTE = RGBColor(0x55, 0x55, 0x55)
FLAGRED = RGBColor(0xB0, 0x00, 0x00)


def title(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(24)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def subtitle(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.color.rgb = GRAYNOTE


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def body(text):
    return doc.add_paragraph(text)


def label(lbl, value):
    p = doc.add_paragraph()
    r = p.add_run(f"{lbl}: ")
    r.bold = True
    r.font.color.rgb = ACCENT
    p.add_run(value)
    return p


def _side(tag, content, is_original, bullet):
    lp = doc.add_paragraph()
    lr = lp.add_run(tag)
    lr.bold = True
    lr.font.size = Pt(9.5)
    lr.font.color.rgb = MUTED if is_original else ACCENT
    if isinstance(content, str):
        content = [content]
    for c in content:
        p = doc.add_paragraph(c, style="List Bullet") if bullet else doc.add_paragraph(c)
        if is_original:
            for rn in p.runs:
                rn.font.color.rgb = MUTED


def chunk(context, a, b, note=None, bullet=False):
    cp = doc.add_paragraph()
    cr = cp.add_run(context)
    cr.bold = True
    cr.font.size = Pt(11)
    cr.font.color.rgb = ACCENT
    _side("A. Original", a, True, bullet)
    _side("B. Optimized", b, False, bullet)
    if note:
        p = doc.add_paragraph()
        r = p.add_run("Why: " + note)
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = GRAYNOTE
    doc.add_paragraph()


def kept(items):
    p = doc.add_paragraph()
    r = p.add_run("Reviewed, kept as-is (already strong): ")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = GRAYNOTE
    for it in items:
        doc.add_paragraph(it, style="List Bullet")
    doc.add_paragraph()


def flag(text):
    p = doc.add_paragraph()
    r = p.add_run("FLAG: " + text)
    r.bold = True
    r.font.color.rgb = FLAGRED
    doc.add_paragraph()


def pagehead(name, path, intent):
    doc.add_page_break()
    doc.add_heading(name, level=1)
    label("Page path", path)
    ip = doc.add_paragraph()
    ir = ip.add_run(intent)
    ir.italic = True
    ir.font.color.rgb = GRAYNOTE
    doc.add_paragraph()


# =====================================================================
# COVER AND FRONT MATTER
# =====================================================================
title("Well and Good Growth: Site v2 Copy Optimization")
subtitle("A and B copy for every page in site-v2. A is the current copy, B is the recommended copy. "
         "Run through the copy-optimization skill stack: cro, angles, copywriting-101, sentence-conscious-writing.")

h2("How to read this deck")
body("Each entry shows the current copy (A) and the recommended replacement (B), with a one line reason. "
     "Chunks that were already strong are listed under Reviewed, kept as-is, so you can see every part was "
     "checked, not just the parts that changed. Red FLAG lines mark open decisions or items that need your input.")

h2("Global decisions applied throughout")
doc.add_paragraph(
    "Brand name. The company is Well and Good Growth. The bare Well and Good is treated as an incomplete name "
    "and replaced wherever it refers to the company. Well and Good Growth Systems is kept as the formal long "
    "form for the About introduction and the legal privacy notice only. The four lines stay as plain service "
    "offerings under the one umbrella, with no separate sub-brands: Websites and Conversion, SEO and Search "
    "Visibility, AEO and AI Visibility, and Agentic Operations.", style="List Bullet")
doc.add_paragraph(
    "Homepage headline. Get found. Get chosen. Operate better. becomes Get found. Get chosen. Get time back. "
    "Each verb now maps to one offering: found is search and AEO, chosen is the website and conversion, time "
    "back is operations. Get time back is the concrete benefit the owner feels, where operate better stayed "
    "abstract.", style="List Bullet")
doc.add_paragraph(
    "Proof. A short proof line is added near the top of the homepage using only the live client and the Jetta "
    "Grove track record. Concept builds are kept on the site but de-emphasized, and are removed from the "
    "homepage proof line.", style="List Bullet")
doc.add_paragraph(
    "Pricing and anchoring. The homepage deliberately carries no dollar figure, to avoid re-anchoring the "
    "practice at the lowest tier. Published pricing stays on the acquisition pages (affordable, one-page, "
    "web design Niagara) and the websites service page, where the entry framing is intentional.", style="List Bullet")
doc.add_paragraph(
    "Consistency. The Launch plan feature list is standardized across the affordable, web design Niagara, and "
    "websites pages to the fullest version, which currently lives on the websites page.", style="List Bullet")
doc.add_paragraph(
    "House rules kept intact: no em-dashes, Claude named alongside ChatGPT, no fabricated statistics or "
    "testimonials, only the real proof assets (Frank Baggetta as live client, Jetta Grove Consulting track "
    "record).", style="List Bullet")

h2("Open decisions to resolve before publishing")
flag("Lock the final name and secure the domain and trademark before publishing. Well and Good Growth is the "
     "working name used throughout this deck. Changing it later is a clean global find and replace.")
flag("Contact email is inconsistent across the site: matt@wellandgoodwebsites.ca on most pages, "
     "baggetta@gmail.com in the privacy notice. Pick one and use it everywhere. This deck flags it but does "
     "not choose for you.")
flag("The contact page links to /services/websites/#preview and /services/websites/#plans. Confirm those "
     "anchors exist on the websites page so the links do not dead-end.")

# =====================================================================
# HOMEPAGE
# =====================================================================
pagehead("Homepage", "index.html",
         "Highest-value page, routes to everything else. Lead angle: relatability at the top (the vendor "
         "juggling pain), belonging near the final call. No price on this page by design.")

chunk("SEO title",
      "Websites, Search Visibility, and AI Operations | Well and Good",
      "Websites, Search Visibility, and AI Operations | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "Well and Good builds websites, search visibility, and AI-powered operations for Niagara and GTA "
      "businesses. Get found on Google and in AI answers, win more customers, and run with less manual work. "
      "Request a free strategy call.",
      "Well and Good Growth builds websites, search visibility, and AI-powered operations for Niagara and GTA "
      "businesses. Get found on Google and in AI answers, win more customers, and run your business with less "
      "manual work. Request a free strategy call.",
      "Brand name, small clarity lift.")

chunk("Hero eyebrow",
      "Well and Good",
      "Well and Good Growth",
      "Brand name.")

chunk("Hero headline (H1)",
      "Get found. Get chosen. Operate better.",
      "Get found. Get chosen. Get time back.",
      "Each verb now maps to one offering (search, conversion, operations). Get time back is concrete where "
      "operate better was abstract. Chosen over Get attention. Get noticed. because those two are near "
      "synonyms and cover only visibility.")

chunk("Hero subheadline",
      "Well and Good builds websites, search visibility, and AI-powered operations for Niagara and GTA "
      "businesses that want to grow without adding overhead. One practice covers the site people land on, the "
      "search and AI results that send them there, and the systems that handle the follow-through.",
      "Well and Good Growth builds websites, search visibility, and AI-powered operations for Niagara and GTA "
      "businesses that want to grow without adding overhead. One practice, one accountable person, covering "
      "the site people land on, the search and AI results that send them there, and the systems that handle "
      "the follow-through.",
      "Brand name, plus the one accountable person differentiator moved up into the subhead.")

chunk("Proof line (new, place under the hero CTAs)",
      "(none on the homepage today; strongest proof is buried on the Work page)",
      "A live client site at frankbaggetta.ca, and a decade of measurable SEO and content results through "
      "Jetta Grove Consulting.",
      "Adds above-the-fold credibility from real assets only. Concept builds omitted here per your call.")

chunk("Section eyebrow",
      "What Well and Good does",
      "What Well and Good Growth does",
      "Brand name.")

chunk("Through-line: section intro",
      "Most businesses buy a website, then a marketing service, then some kind of automation, from three "
      "different vendors who never talk to each other. Well and Good runs all three as one connected system, "
      "so the parts reinforce each other instead of working around each other.",
      "Most businesses buy a website, then a marketing service, then some kind of automation, from three "
      "different vendors who never talk to each other. Well and Good Growth runs all three as one connected "
      "system, so the parts strengthen each other instead of quietly working against each other.",
      "Brand name, plus a sharper closing verb pair on read-aloud.")

chunk("Why one practice: body",
      "A website that ranks well but cannot handle the leads it generates is only half a solution. Search "
      "visibility that fills your inbox faster than you can reply creates a new problem. Well and Good "
      "connects the site, the search and AI presence, and the operations behind them, so growth in one area "
      "does not break another. You work with one person who understands the whole picture.",
      "A website that ranks well but cannot handle the leads it generates is only half a solution. Search "
      "visibility that fills your inbox faster than you can answer just moves the bottleneck. Well and Good "
      "Growth connects the site, the search and AI presence, and the operations behind them, so growth in one "
      "area does not break another. You work with one person who understands the whole picture.",
      "Brand name, plus a tighter loss-aversion beat (moves the bottleneck) that reads better aloud.")

chunk("Selected work: section intro",
      "A live client site, two concept builds, and a decade of prior consultancy results in search and "
      "content.",
      "A live client site and a decade of measurable search and content results, plus concept builds that "
      "show how specific business types can be structured.",
      "Leads with the hardest proof and moves concept builds into a supporting clause, per keep but "
      "de-emphasize.")

chunk("Final CTA: heading",
      "Not sure which part you need first?",
      "Not sure which part you need first? Most owners we start with are not either.",
      "Adds a belonging beat right before the ask, without inventing anything.")

chunk("Final CTA: body",
      "That is what the strategy call is for. It is a short conversation about what is not working, what a "
      "useful first step could be, and whether Well and Good is the right fit. If a simple website is "
      "genuinely all you need, we will point you straight to it.",
      "That is what the strategy call is for. It is a short conversation about what is not working, what a "
      "useful first step could be, and whether Well and Good Growth is the right fit. If a simple website is "
      "genuinely all you need, we will point you straight to it.",
      "Brand name.")

kept([
    "Hero CTAs: Request a free strategy call (primary) and See the services (secondary). Clear hierarchy; no "
    "price CTA added here on purpose.",
    "The three capability cards (Websites and conversion, SEO and AI visibility, Agentic operations). Benefit "
    "led and honest, including human approval stays in place.",
    "Who this is for list. Specific and self-qualifying.",
    "Frank Baggetta and the Jetta Grove consultancy cards.",
    "Micro line: Free. No obligation. Usually a reply within one business day.",
])

# =====================================================================
# WEBSITES SERVICE PAGE
# =====================================================================
pagehead("Websites and Conversion", "services/websites/index.html",
         "The accessible entry offer and the home of published pricing. This page holds the canonical Launch "
         "plan feature list that the acquisition pages are standardized to.")

chunk("SEO title",
      "Website Design and Conversion for Local Business | Well and Good",
      "Website Design and Conversion for Local Business | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "Custom websites and conversion paths for Niagara and GTA businesses, built to be found and to turn "
      "visitors into calls, bookings, and quotes. Website and growth plans from $99/month, published pricing. "
      "Get a free website preview.",
      "Custom websites and conversion paths for Niagara and GTA businesses, built to be found and to turn "
      "visitors into calls, bookings, and quotes. Website and growth plans from $99 per month, published "
      "pricing. Get a free website preview from Well and Good Growth.",
      "Brand name, and $99 per month reads more naturally than $99/month in a sentence.")

chunk("Hero subheadline",
      "Your website is the foundation everything else runs on. Well and Good builds custom sites for Niagara "
      "and GTA businesses that load fast, read well on a phone, and make the next step obvious. Every build "
      "ships optimized for local search, so customers on Google and in AI answers can find you. Published "
      "pricing, no lock-in.",
      "Your website is the foundation everything else runs on. Well and Good Growth builds custom sites for "
      "Niagara and GTA businesses that load fast, read well on a phone, and make the next step obvious. Every "
      "build ships optimized for local search, so customers on Google and in AI answers can find you. "
      "Published pricing, no lock-in.",
      "Brand name.")

chunk("FAQ: Is this better than Wix or Squarespace?",
      "Those builders work if you want to do the work yourself. Well and Good builds the site for you, writes "
      "the copy, sets up the local search signals, and maintains it, so you are running your business instead "
      "of editing a website.",
      "Those builders work if you want to do the work yourself. Well and Good Growth builds the site for you, "
      "writes the copy, sets up the local search signals, and maintains it, so you are running your business "
      "instead of editing a website.",
      "Brand name.")

kept([
    "H1: A website built to be found and to convert. Clear and benefit led.",
    "Price band ($99/mo, or a one-time build from $297). Published pricing is a differentiator here and stays.",
    "What is included section and its four cards. Concrete and well ordered.",
    "Launch, Grow, Dominate plan cards. This page holds the canonical feature lists.",
    "The one-time builds paragraph and the Care add-on.",
    "Remaining FAQs and the free preview final CTA. Strong as written.",
])

flag("Confirm the #preview and #plans anchors exist on this page. The contact page and other pages link to "
     "them directly.")

# =====================================================================
# AFFORDABLE WEBSITE DESIGN
# =====================================================================
pagehead("Affordable Website Design", "affordable-website-design/index.html",
         "Acquisition page targeting affordable web design search intent. Pricing stays. Main fix here is the "
         "Launch plan feature mismatch against the websites page.")

chunk("SEO title",
      "Affordable Website Design in Niagara | Well and Good",
      "Affordable Website Design in Niagara | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "Affordable custom websites for Niagara and GTA businesses, plus the SEO, Google Business Profile, and "
      "social that get you found on Google and in AI answers. From $99/month, published pricing, no lock-in.",
      "Affordable custom websites for Niagara and GTA businesses, plus the SEO, Google Business Profile, and "
      "social that get you found on Google and in AI answers from ChatGPT and Claude. From $99 per month, "
      "published pricing, no lock-in.",
      "Brand consistency and pairs Claude with ChatGPT in the meta, matching the house rule.")

chunk("Launch plan: feature list (fix the cross-page mismatch)",
      ["One-page website plus hosting and care",
       "Google Business Profile management",
       "Local SEO foundation plus citations",
       "4 social posts per month plus a monthly report"],
      ["One-page website, mobile, with tap-to-call",
       "Google Business Profile setup and management",
       "Local SEO foundation and core citations",
       "Hosting, SSL, backups, and monitoring",
       "4 social posts per month on 1 platform, with a monthly results email"],
      "Standardizes Launch to the canonical websites-page list, so the same plan reads the same everywhere.",
      bullet=True)

chunk("FAQ: How much does an affordable website really cost?",
      "A one-time build starts at $297 for Express. If you want ongoing growth alongside the site, the Launch "
      "plan is $99 per month plus a $497 setup. Every figure is published on this page, so there is no "
      "request-a-quote step.",
      "A one-time build starts at $297 for Express. If you want ongoing growth alongside the site, the Launch "
      "plan is $99 per month plus a $497 setup. Every figure is published on this page, so there is no "
      "request-a-quote step, and no surprise invoice later.",
      "Small trust addition on the transparency angle. Facts unchanged.")

kept([
    "H1: A custom website, and the growth that fills it. Strong and on strategy.",
    "Affordable should still mean custom section. Answers the cheap template objection well.",
    "Grow and Dominate plan cards, and the one-time builds paragraph.",
    "Remaining FAQs (professional look, contract) and the free preview final CTA.",
])

# =====================================================================
# WEB DESIGN NIAGARA
# =====================================================================
pagehead("Web Design Niagara", "web-design-niagara/index.html",
         "Local acquisition page. Same Launch plan fix as the affordable page, plus one dangling sentence.")

chunk("SEO title",
      "Web Design Niagara: Websites, Local SEO, and Growth | Well and Good",
      "Web Design Niagara: Websites, Local SEO, and Growth | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "Web design and growth for Niagara businesses. Get found on Google and in AI answers with custom "
      "websites plus local SEO, Google Business Profile, and social. From $99/month, published pricing, no "
      "lock-in.",
      "Web design and growth for Niagara businesses. Get found on Google and in AI answers from ChatGPT and "
      "Claude with custom websites plus local SEO, Google Business Profile, and social. From $99 per month, "
      "published pricing, no lock-in.",
      "Pairs Claude with ChatGPT in the meta, brand consistency.")

chunk("Local signals card: body",
      "That means clear service language, structured pages, local signals, fast load times, and useful "
      "content that helps customers understand what you offer.",
      "Clear service language, structured pages, local signals, fast load times, and useful content that "
      "helps customers understand what you offer.",
      "Drops the dangling That means, which had no antecedent as a card opener.")

chunk("Launch plan: feature list (fix the cross-page mismatch)",
      ["One-page website plus hosting and care",
       "Google Business Profile management",
       "Local SEO foundation plus citations",
       "4 social posts per month plus a monthly report"],
      ["One-page website, mobile, with tap-to-call",
       "Google Business Profile setup and management",
       "Local SEO foundation and core citations",
       "Hosting, SSL, backups, and monitoring",
       "4 social posts per month on 1 platform, with a monthly results email"],
      "Same standardization as the affordable page.",
      bullet=True)

chunk("FAQ: How much does web design in Niagara cost with Well and Good?",
      "A one-time build starts at $297. A website plus ongoing local growth starts at $99 per month with a "
      "$497 setup. Every price is published here, which most Niagara agencies do not do.",
      "A one-time build starts at $297. A website plus ongoing local growth starts at $99 per month with a "
      "$497 setup. Every price is published here, which most Niagara agencies do not do. (Question heading "
      "updates to: How much does web design in Niagara cost with Well and Good Growth?)",
      "Brand name in the question heading. Answer facts unchanged.")

kept([
    "H1: Get your Niagara business found and chosen. Strong local hook.",
    "Built around how Niagara customers choose section (Local proof, Mobile actions).",
    "Grow and Dominate plan cards and the one-time builds line.",
    "Remaining FAQs and the free preview final CTA.",
])

# =====================================================================
# SERVICES INDEX
# =====================================================================
pagehead("Services", "services/index.html",
         "Hub page for the four offerings. Mostly a brand-name pass; the structure and honesty already fit "
         "the repositioning.")

chunk("SEO title",
      "Services: Websites, SEO, AEO, and AI Operations | Well and Good",
      "Services: Websites, SEO, AEO, and AI Operations | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "Well and Good offers four connected services for Niagara and GTA businesses: websites and conversion, "
      "SEO and search visibility, AEO and AI visibility, and agentic operations. See how they fit together.",
      "Well and Good Growth offers four connected services for Niagara and GTA businesses: websites and "
      "conversion, SEO and search visibility, AEO and AI visibility, and agentic operations. See how they fit "
      "together.",
      "Brand name.")

chunk("Section intro",
      "Well and Good covers the foundation your business runs on, the visibility that brings people to it, "
      "and the operations that handle what happens next. You can start with one and add others as it makes "
      "sense. Each service page below explains what it is, how it works, and what to expect.",
      "Well and Good Growth covers the foundation your business runs on, the visibility that brings people to "
      "it, and the operations that handle what happens next. You can start with one and add others as it "
      "makes sense. Each service page below explains what it is, how it works, and what to expect.",
      "Brand name.")

kept([
    "H1: Four services, designed to reinforce each other.",
    "The four service blocks, including the honest No method guarantees a citation line under AEO.",
    "How engagements are scoped section and the final CTA.",
])

# =====================================================================
# SEO SERVICE PAGE
# =====================================================================
pagehead("SEO and Search Visibility", "services/seo/index.html",
         "Authority service page. Brand-name pass; the definition-first structure already suits AEO.")

chunk("SEO title",
      "SEO and Local Search Visibility | Well and Good",
      "SEO and Local Search Visibility | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "SEO and local search visibility for Niagara and GTA businesses. Technical SEO, on-page work, local "
      "SEO, Google Business Profile, schema, and content systems that help customers find you on Google. "
      "Request a free strategy call.",
      "SEO and local search visibility for Niagara and GTA businesses from Well and Good Growth. Technical "
      "SEO, on-page work, local SEO, Google Business Profile, schema, and content systems that help customers "
      "find you on Google. Request a free strategy call.",
      "Brand name.")

chunk("Hero subheadline",
      "SEO is the work that helps your business show up when people search for what you do. Well and Good "
      "handles the technical foundation, the on-page details, local signals, and the content that answers "
      "real customer questions. The aim is steady, qualified visibility in the searches that lead to calls, "
      "bookings, and quotes.",
      "SEO is the work that helps your business show up when people search for what you do. Well and Good "
      "Growth handles the technical foundation, the on-page details, local signals, and the content that "
      "answers real customer questions. The aim is steady, qualified visibility in the searches that lead to "
      "calls, bookings, and quotes.",
      "Brand name.")

chunk("Section heading",
      "What SEO with Well and Good covers",
      "What SEO with Well and Good Growth covers",
      "Brand name.")

kept([
    "H1: Get found when customers search Google.",
    "The five work cards (Technical, On-page, Local, Content systems, Schema).",
    "SEO and AEO are related, not the same section.",
    "What SEO can and cannot promise, including the honest track record framing tied to Jetta Grove.",
    "All three FAQs and the final CTA.",
])

# =====================================================================
# AEO SERVICE PAGE
# =====================================================================
pagehead("AEO and AI Visibility", "services/aeo/index.html",
         "The clearest expression of the honesty positioning. Brand-name pass only; the copy already pairs "
         "ChatGPT and Claude and refuses to over-promise.")

chunk("SEO title",
      "AEO and AI Visibility: Get Cited by ChatGPT and Claude | Well and Good",
      "AEO and AI Visibility: Get Cited by ChatGPT and Claude | Well and Good Growth",
      "Brand name. ChatGPT and Claude already paired.")

chunk("Meta description",
      "Answer engine optimization for Niagara and GTA businesses. Structure your website and facts so AI "
      "answer engines like ChatGPT, Claude, and Perplexity can find and cite your business. Honest about "
      "limits. Request a free strategy call.",
      "Answer engine optimization for Niagara and GTA businesses from Well and Good Growth. Structure your "
      "website and facts so AI answer engines like ChatGPT, Claude, and Perplexity can find and cite your "
      "business. Honest about limits. Request a free strategy call.",
      "Brand name.")

chunk("Hero subheadline",
      "More people now ask ChatGPT, Claude, Perplexity, and Google's AI answers for recommendations instead "
      "of scrolling a page of blue links. AEO is the work that makes your business easy for those systems to "
      "find, understand, and cite. Well and Good does this work directly. No method guarantees a citation.",
      "More people now ask ChatGPT, Claude, Perplexity, and Google's AI answers for recommendations instead "
      "of scrolling a page of blue links. AEO is the work that makes your business easy for those systems to "
      "find, understand, and cite. Well and Good Growth does this work directly. No method guarantees a "
      "citation.",
      "Brand name. The honest closing line stays.")

chunk("Section heading",
      "What AEO with Well and Good covers",
      "What AEO with Well and Good Growth covers",
      "Brand name.")

chunk("Honesty section: body (brand mention)",
      "No one can guarantee that an AI system will cite your business, and Well and Good will never claim "
      "otherwise.",
      "No one can guarantee that an AI system will cite your business, and Well and Good Growth will never "
      "claim otherwise.",
      "Brand name. This whole section is a strength and is otherwise untouched.")

kept([
    "H1: Show up when customers ask AI, not just Google.",
    "What is AEO and How is AEO different from SEO explainer blocks.",
    "The five work cards, including the honest llms.txt card.",
    "What llms.txt does and does not do.",
    "All three FAQs, which correctly refuse to promise citations. Final CTA.",
])

# =====================================================================
# AUTOMATION SERVICE PAGE
# =====================================================================
pagehead("Agentic Operations", "services/automation/index.html",
         "The highest-value offer. Brand-name pass; the human-approval honesty is central and preserved.")

chunk("SEO title",
      "Agentic Operations and AI Automation for Business | Well and Good",
      "Agentic Operations and AI Automation for Business | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "Agentic operations for Niagara and GTA businesses. AI workflows for lead intake, research, reporting, "
      "inboxes, documents, and meeting follow-up, with human approval where judgment and money are involved. "
      "Request a free strategy call.",
      "Agentic operations for Niagara and GTA businesses from Well and Good Growth. AI workflows for lead "
      "intake, research, reporting, inboxes, documents, and meeting follow-up, with human approval where "
      "judgment and money are involved. Request a free strategy call.",
      "Brand name.")

chunk("Hero subheadline",
      "Agentic operations is the highest-value work Well and Good does. It means designing AI workflows that "
      "handle the repetitive tasks eating your week, lead intake, research, reporting, inbox and document "
      "handling, meeting follow-up, so your time goes to the work only you can do. Human approval stays in "
      "place wherever judgment, money, or customer communication is involved.",
      "Agentic operations is the highest-value work Well and Good Growth does. It means designing AI "
      "workflows that handle the repetitive tasks eating your week, lead intake, research, reporting, inbox "
      "and document handling, meeting follow-up, so your time goes to the work only you can do. Human "
      "approval stays in place wherever judgment, money, or customer communication is involved.",
      "Brand name.")

chunk("Definition and design mentions (brand name, three spots)",
      "Well and Good designs these workflows with human approval at the points where judgment, money, or "
      "communication matter. ... These are the workflows Well and Good designs. ... Well and Good designs "
      "agents to prepare the work and pause for approval at those points.",
      "Well and Good Growth designs these workflows with human approval at the points where judgment, money, "
      "or communication matter. ... These are the workflows Well and Good Growth designs. ... Well and Good "
      "Growth designs agents to prepare the work and pause for approval at those points.",
      "Brand name in the three inline references. Wording otherwise unchanged.")

chunk("What we run today: body (brand mention)",
      "Well and Good's automation work is demonstrated through its own internal systems.",
      "Well and Good Growth's automation work is demonstrated through its own internal systems.",
      "Brand name.")

kept([
    "H1: Automate the repetitive work, keep control of the judgment. Strong tension in the headline.",
    "The seven workflow cards and the human oversight principle section.",
    "You can start with a single workflow section.",
    "All four FAQs, including What should not be automated. Final CTA.",
])

# =====================================================================
# ABOUT
# =====================================================================
pagehead("About", "about/index.html",
         "The one place the formal long form Well and Good Growth Systems is used on first mention, then the "
         "short name. Strong page; brand-name pass.")

chunk("SEO title",
      "About Well and Good and Matthew Baggetta | Niagara",
      "About Well and Good Growth and Matthew Baggetta | Niagara",
      "Brand name.")

chunk("Meta description",
      "Well and Good is a focused practice run by Matthew Baggetta, serving Niagara and GTA businesses with "
      "websites, search visibility, and AI-powered operations. Read the story and the experience.",
      "Well and Good Growth is a focused practice run by Matthew Baggetta, serving Niagara and GTA businesses "
      "with websites, search visibility, and AI-powered operations. Read the story and the experience.",
      "Brand name.")

chunk("Intro body (first mention uses the formal long form)",
      "Well and Good is a focused practice run by Matthew Baggetta in Welland, serving Niagara and the GTA. "
      "When you work with Well and Good, you work with the person doing the work, not a rotating account "
      "team. That means one line of accountability across your website, your search presence, and the "
      "systems that run behind them.",
      "Well and Good Growth Systems is a focused practice run by Matthew Baggetta in Welland, serving Niagara "
      "and the GTA. When you work with Well and Good Growth, you work with the person doing the work, not a "
      "rotating account team. That means one line of accountability across your website, your search "
      "presence, and the systems that run behind them.",
      "Formal long form on the first mention, short name after. This is the only page that introduces the "
      "full name.")

chunk("Why this practice exists: body (brand mention)",
      "Well and Good exists to run those three together as one system.",
      "Well and Good Growth exists to run those three together as one system.",
      "Brand name.")

kept([
    "H1: One person, accountable for the whole picture. This is the differentiator and it is well stated.",
    "The experience behind the work section (Jetta Grove, honest and attributed).",
    "How I work section and the four values cards.",
    "Final CTA.",
])

# =====================================================================
# WORK
# =====================================================================
pagehead("Work", "work/index.html",
         "The proof page. Concept builds stay here, clearly labelled. Real numbers from Jetta Grove stay, "
         "with attribution. Brand-name pass.")

chunk("SEO title",
      "Work: Client Sites, Concept Builds, and Demonstrations | Well and Good",
      "Work: Client Sites, Concept Builds, and Demonstrations | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "Selected work from Well and Good: a live client website, labelled concept builds, prior consultancy "
      "results, and internal automation demonstrations. Honest about what is client work and what is not.",
      "Selected work from Well and Good Growth: a live client website, labelled concept builds, prior "
      "consultancy results, and internal automation demonstrations. Honest about what is client work and "
      "what is not.",
      "Brand name.")

chunk("Intro body",
      "This page shows what Well and Good has actually built, with labels that say what each piece is. A live "
      "client site is a live client site. A concept build is marked as a concept. Prior consultancy results "
      "are attributed to that prior work, not to Well and Good. Internal automation is shown as internal "
      "systems.",
      "This page shows what Well and Good Growth has actually built, with labels that say what each piece is. "
      "A live client site is a live client site. A concept build is marked as a concept. Prior consultancy "
      "results are attributed to that prior work, not to Well and Good Growth. Internal automation is shown "
      "as internal systems.",
      "Brand name in both references.")

chunk("Automation demonstrations: body (brand mentions)",
      "Well and Good runs its own AI workflows, shown here as live demonstrations of the operations work "
      "offered on this site. ... An internal system that checks how Well and Good appears across Google, "
      "ChatGPT, Claude, Perplexity, and Gemini.",
      "Well and Good Growth runs its own AI workflows, shown here as live demonstrations of the operations "
      "work offered on this site. ... An internal system that checks how Well and Good Growth appears across "
      "Google, ChatGPT, Claude, Perplexity, and Gemini.",
      "Brand name. ChatGPT and Claude already paired.")

kept([
    "H1: Real work, and the honest labelling throughout the page.",
    "Frank Baggetta live client block.",
    "Concept build blocks (Evelyn's Sandwich Factory, JK Motors), kept and clearly labelled per your call.",
    "Jetta Grove stat block (163 keywords, Page 1, 4.34M impressions, referral lift) with the attribution "
    "line. Real numbers, do not touch.",
    "The three internal demonstration cards and final CTA.",
])

# =====================================================================
# ENGAGEMENTS
# =====================================================================
pagehead("How It Works", "engagements/index.html",
         "The pricing-conversation page that replaces hard anchors with scope. Brand-name pass.")

chunk("SEO title",
      "How Engagements Work: Scoping, Discovery, and Pricing | Well and Good",
      "How Engagements Work: Scoping, Discovery, and Pricing | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "How Well and Good scopes website projects, search programs, and operations systems for Niagara and GTA "
      "businesses. What determines cost, what discovery includes, and how to start with one step. Request a "
      "free strategy call.",
      "How Well and Good Growth scopes website projects, search programs, and operations systems for Niagara "
      "and GTA businesses. What determines cost, what discovery includes, and how to start with one step. "
      "Request a free strategy call.",
      "Brand name.")

chunk("Intro body",
      "Well and Good keeps the process simple. Website projects, search programs, and operations systems are "
      "each scoped to the work involved, so you pay for what your situation needs rather than a fixed package "
      "that does not fit. This page explains the three kinds of engagement, what determines cost, and how to "
      "begin.",
      "Well and Good Growth keeps the process simple. Website projects, search programs, and operations "
      "systems are each scoped to the work involved, so you pay for what your situation needs rather than a "
      "fixed package that does not fit. This page explains the three kinds of engagement, what determines "
      "cost, and how to begin.",
      "Brand name.")

kept([
    "H1: How working together actually works.",
    "The three engagement type cards and the What determines the cost section.",
    "What discovery includes.",
    "All five FAQs, including the honest When does automation produce a worthwhile return. Final CTA.",
])

# =====================================================================
# ONE-PAGE WEBSITES
# =====================================================================
pagehead("One-Page Websites", "one-page-websites/index.html",
         "Acquisition page for one-page website search intent. Already clean; brand and meta pass.")

chunk("SEO title",
      "One-Page Websites for Local Businesses | Well and Good",
      "One-Page Websites for Local Businesses | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "Custom one-page websites for local businesses in Niagara and the GTA, built to be found on Google and "
      "in AI answers and to turn visitors into calls and bookings. Available bundled with growth from "
      "$99/month.",
      "Custom one-page websites for local businesses in Niagara and the GTA, built to be found on Google and "
      "in AI answers from ChatGPT and Claude and to turn visitors into calls and bookings. Bundled with "
      "growth from $99 per month.",
      "Pairs Claude with ChatGPT, brand consistency.")

kept([
    "H1: One-page websites that make the next step obvious.",
    "When a one-page website works well section and its three cards.",
    "What the page needs to include (Trust, Search clarity, Action bullet groups).",
    "Pricing section, all three FAQs, and the free preview final CTA.",
])

# =====================================================================
# CONTACT
# =====================================================================
pagehead("Contact", "contact/index.html",
         "Conversion page. Brand-name pass. Flags the email inconsistency and the anchor links.")

chunk("SEO title",
      "Contact Well and Good | Book a Free Strategy Call",
      "Contact Well and Good Growth | Book a Free Strategy Call",
      "Brand name.")

chunk("Meta description",
      "Request a free strategy call with Well and Good. A short, direct conversation about your website, "
      "search visibility, or operations, and whether Well and Good is the right fit. Serving Niagara and the "
      "GTA.",
      "Request a free strategy call with Well and Good Growth. A short, direct conversation about your "
      "website, search visibility, or operations, and whether Well and Good Growth is the right fit. Serving "
      "Niagara and the GTA.",
      "Brand name.")

chunk("Intro body",
      "The strategy call is a short, plain conversation, not a sales pitch. We talk about what is not "
      "working, what a useful first step could be, and whether Well and Good is the right fit. It is free and "
      "carries no obligation. If Well and Good is not the right choice for your situation, I will tell you.",
      "The strategy call is a short, plain conversation, not a sales pitch. We talk about what is not "
      "working, what a useful first step could be, and whether Well and Good Growth is the right fit. It is "
      "free and carries no obligation. If Well and Good Growth is not the right choice for your situation, I "
      "will tell you.",
      "Brand name in both references.")

kept([
    "H1: Request a free strategy call.",
    "What the call covers section and the Just want a website alternative path.",
    "The needs selector form and its options (Website, SEO, AEO, Operations automation, Not sure yet).",
])

flag("The email inconsistency surfaces here too. This page uses matt@wellandgoodwebsites.ca, which is fine, "
     "but the privacy notice uses baggetta@gmail.com. Align them.")

# =====================================================================
# THANK YOU
# =====================================================================
pagehead("Thank You", "thank-you/index.html",
         "Post-submit page. Body does not name the company, so this is a title and meta pass only.")

chunk("SEO title",
      "Thank You | Well and Good",
      "Thank You | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "Thanks for reaching out to Well and Good. Your message is in, and I will reply personally, usually "
      "within one business day.",
      "Thanks for reaching out to Well and Good Growth. Your message is in, and I will reply personally, "
      "usually within one business day.",
      "Brand name.")

kept([
    "H1: Got it. Your message is in. Warm and human.",
    "What happens next and While you wait sections. The three onward links are well chosen.",
])

# =====================================================================
# 404
# =====================================================================
pagehead("404 Not Found", "404.html",
         "Utility page. Title, meta, and the one brand reference in the meta description.")

chunk("SEO title",
      "Page Not Found | Well and Good",
      "Page Not Found | Well and Good Growth",
      "Brand name.")

chunk("Meta description",
      "That page could not be found. Head back to the homepage or explore Well and Good's services, work, and "
      "how engagements work.",
      "That page could not be found. Head back to the homepage or explore Well and Good Growth's services, "
      "work, and how engagements work.",
      "Brand name.")

kept([
    "H1: That page is not here, and the reassuring body copy (Nothing is broken on your end).",
    "The email fallback line.",
])

# =====================================================================
# PRIVACY
# =====================================================================
pagehead("Privacy Notice", "privacy/index.html",
         "Legal page. Uses the formal long form for the operating entity. Otherwise left as legal text.")

chunk("SEO title",
      "Privacy Notice | Well and Good Websites",
      "Privacy Notice | Well and Good Growth",
      "Brand name. The old title still carried the Websites-only identity.")

chunk("Meta description",
      "Read how Well and Good, operated by Matthew Baggetta in Ontario, handles the details you send through "
      "our forms and how you can access or remove them.",
      "Read how Well and Good Growth, operated by Matthew Baggetta in Ontario, handles the details you send "
      "through our forms and how you can access or remove them.",
      "Brand name.")

chunk("Opening line (operating entity, formal long form)",
      "Well and Good is operated by Matthew Baggetta in Ontario, Canada. This notice explains what "
      "information we collect through the website, how we use it, and the choices you have.",
      "Well and Good Growth Systems is operated by Matthew Baggetta in Ontario, Canada. This notice explains "
      "what information we collect through the website, how we use it, and the choices you have.",
      "Formal long form names the legal operator. The rest of the notice is left as written.")

flag("Privacy contact is baggetta@gmail.com while the rest of the site uses matt@wellandgoodwebsites.ca. "
     "Decide on one address. If the domain changes with the rename, update it here too.")

kept([
    "All substantive privacy clauses (Information we collect, How we use, Form delivery and hosting via "
    "FormSubmit and Vercel, Security, Your choices). Legal text, left intact except the operating name.",
])

# =====================================================================
out = str(pathlib.Path(__file__).parent / "Well-and-Good-Growth-Site-v2-Copy-Optimization.docx")
doc.save(out)
print("Saved:", out)
